"""Regressions for preserving user files and validating metadata mutations."""
from pathlib import Path
from zipfile import ZipFile
import sys

import pytest

from core.models import MetadataField
import core.path_security as security


def test_atomic_write_preserves_existing_neighbor(tmp_path):
    target = tmp_path / "original.txt"
    neighbor = tmp_path / "original.ml_tmp.txt"
    target.write_text("original")
    neighbor.write_text("user data")
    security.secure_atomic_write(target, lambda p: p.write_text("changed"))
    assert target.read_text() == "changed"
    assert neighbor.read_text() == "user data"


def test_atomic_write_detects_concurrent_change(tmp_path):
    target = tmp_path / "original.txt"
    target.write_text("original")
    def writer(p):
        p.write_text("new")
        target.write_text("external change")
    with pytest.raises(Exception, match="changed"):
        security.secure_atomic_write(target, writer)
    assert target.read_text() == "external change"


def test_expand_home(tmp_path, monkeypatch):
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    monkeypatch.setenv("HOME", str(tmp_path))
    target = tmp_path / "file.txt"
    target.write_text("ok")
    assert security.validate_file_path("~/file.txt") == target.resolve()


def test_pdf_write_preserves_other_metadata(tmp_path):
    from pypdf import PdfReader, PdfWriter
    from core.handlers.pdf_handler import PdfHandler
    target = tmp_path / "file.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_metadata({"/Title": "old", "/Author": "keep", "/Custom": "also keep"})
    writer.write(target)
    PdfHandler().write(target, [MetadataField("pdf:Title", "Title", "new")])
    meta = PdfReader(target).metadata
    assert meta.title == "new"
    assert meta.author == "keep"
    assert meta["/Custom"] == "also keep"


def test_office_delete_and_date_roundtrip(tmp_path):
    from docx import Document
    from core.handlers.office_handler import OfficeHandler
    target = tmp_path / "file.docx"
    doc = Document()
    doc.core_properties.title = "remove me"
    doc.save(target)
    handler = OfficeHandler()
    handler.delete(target, ["office:title"])
    assert Document(target).core_properties.title == ""
    handler.write(target, [MetadataField("office:created", "Created", "2022-01-01T00:00:00Z")])
    assert Document(target).core_properties.created.year == 2022


def test_office_preserves_all_other_zip_members(tmp_path):
    from openpyxl import Workbook
    from core.handlers.office_handler import OfficeHandler
    target = tmp_path / "file.xlsm"
    Workbook().save(target)
    with ZipFile(target, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro content")
        archive.writestr("customXml/item1.xml", b"<preserve/>")
    with ZipFile(target) as archive:
        before = {i.filename: archive.read(i) for i in archive.infolist()}
    OfficeHandler().write(target, [MetadataField("office:title", "Title", "new")])
    with ZipFile(target) as archive:
        after = {i.filename: archive.read(i) for i in archive.infolist()}
    assert set(after) == set(before)
    assert all(after[k] == v for k, v in before.items() if k != "docProps/core.xml")


@pytest.mark.parametrize("extension", ["png", "gif", "bmp", "tiff", "webp"])
def test_image_capabilities_are_honest(tmp_path, extension):
    from PIL import Image
    from core.handlers.image_handler import ImageHandler
    target = tmp_path / ("file." + extension)
    Image.new("RGB", (2, 2)).save(target)
    record = ImageHandler().read(target)
    assert not record.supports_write
    assert not record.supports_delete
    assert all(not f.editable and not f.deletable for f in record.fields)


def test_exif_rational_roundtrip(tmp_path):
    import piexif
    from PIL import Image
    from core.handlers.image_handler import ImageHandler
    target = tmp_path / "file.jpg"
    Image.new("RGB", (2, 2)).save(target, exif=piexif.dump({"0th": {piexif.ImageIFD.XResolution: (96, 1)}}))
    handler = ImageHandler()
    handler.write(target, [MetadataField("exif:XResolution", "X Resolution", "120/1")])
    assert piexif.load(str(target))["0th"][piexif.ImageIFD.XResolution] == (120, 1)


@pytest.mark.parametrize("value,tag_type", [("1/0", 5), (str(2**40), 5), (-1, 5), (65536, 3), (-1, 1)])
def test_exif_rejects_invalid_numeric_values(value, tag_type):
    from core.handlers.image_handler import _encode_exif_value
    with pytest.raises(ValueError):
        _encode_exif_value(value, tag_type)


@pytest.mark.parametrize("a,b", [("Alice", "alice"), (" text", "text"), (b"A", b"a")])
def test_diff_preserves_case_and_whitespace(a, b):
    from core.diff import _values_equal
    assert not _values_equal(a, b)


@pytest.mark.asyncio
async def test_api_rejects_missing_field_values(client, tmp_path):
    target = tmp_path / "file.txt"
    target.write_text("ok")
    response = await client.post("/write", json={"path": str(target), "fields": [{"key": "fs:mtime"}]})
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_api_routes_filesystem_fields_on_pdf(client, tmp_path):
    from pypdf import PdfWriter
    target = tmp_path / "file.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(target)
    response = await client.post("/write", json={"path": str(target), "fields": [{"key": "fs:mtime", "value": "2020-01-01T00:00:00Z"}]})
    assert response.status_code == 200
    assert target.stat().st_mtime == pytest.approx(1577836800, abs=1)


@pytest.mark.asyncio
async def test_api_rejects_read_only_and_unknown_fields(client, tmp_path):
    target = tmp_path / "file.txt"
    target.write_text("ok")
    for key in ("fs:size", "unknown:value"):
        response = await client.post("/write", json={"path": str(target), "fields": [{"key": key, "value": "x"}]})
        assert response.status_code == 422


@pytest.mark.parametrize("name,extension", [("audio", "mp3"), ("video", "mp4")])
@pytest.mark.parametrize("operation", ["write", "delete"])
def test_media_save_failure_preserves_original(tmp_path, monkeypatch, name, extension, operation):
    import importlib
    module = importlib.import_module(f"core.handlers.{name}_handler")
    handler = getattr(module, name.title() + "Handler")()
    target = tmp_path / ("media." + extension)
    target.write_bytes(b"original content")
    method = "_write_tags" if operation == "write" else "_delete_tags"
    def failing_save(path, values):
        path.write_bytes(b"partial write")
        raise OSError("simulated disk failure")
    monkeypatch.setattr(handler, method, failing_save)
    with pytest.raises(OSError, match="simulated"):
        getattr(handler, operation)(target, [])
    assert target.read_bytes() == b"original content"
    assert list(tmp_path.iterdir()) == [target]


def test_atomic_replace_failure_preserves_original_and_cleans(tmp_path, monkeypatch):
    target = tmp_path / "file.txt"
    target.write_text("original")
    def fail_replace(*args):
        raise PermissionError("locked file")
    monkeypatch.setattr(security, "_replace_file", fail_replace)
    with pytest.raises(PermissionError, match="locked"):
        security.secure_atomic_write(target, lambda p: p.write_text("new"))
    assert target.read_text() == "original"
    assert list(tmp_path.iterdir()) == [target]


def test_pdf_delete_preserves_other_metadata(tmp_path):
    from pypdf import PdfReader, PdfWriter
    from core.handlers.pdf_handler import PdfHandler
    target = tmp_path / "file.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_metadata({"/Title": "remove", "/Author": "keep"})
    writer.write(target)
    PdfHandler().delete(target, ["pdf:Title"])
    reader = PdfReader(target)
    assert "/Title" not in reader.metadata
    assert reader.metadata.author == "keep"
    assert len(reader.pages) == 1


@pytest.mark.asyncio
async def test_api_mixed_write_rolls_back_on_invalid_datetime(client, tmp_path):
    from pypdf import PdfWriter
    target = tmp_path / "file.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_metadata({"/Title": "original"})
    writer.write(target)
    before = target.read_bytes()
    response = await client.post("/write", json={"path": str(target), "fields": [
        {"key": "pdf:Title", "value": "changed"}, {"key": "fs:mtime", "value": "invalid"},
    ]})
    assert response.status_code == 422
    assert target.read_bytes() == before


@pytest.mark.asyncio
async def test_api_delete_rejects_immutable_field(client, tmp_path):
    target = tmp_path / "file.txt"
    target.write_text("ok")
    response = await client.post("/delete", json={"path": str(target), "keys": ["fs:size"]})
    assert response.status_code == 422


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", ["jpeg", "pdf", "docx", "mp3"])
async def test_api_delete_then_restore_uses_server_permissions(client, tmp_path, kind):
    target = tmp_path / ("file." + kind)
    if kind == "jpeg":
        import piexif
        from PIL import Image
        Image.new("RGB", (2, 2)).save(target, exif=piexif.dump({"0th": {piexif.ImageIFD.Make: b"Original"}}))
        key, value = "exif:Make", "Original"
    elif kind == "pdf":
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.add_metadata({"/Title": "Original"})
        writer.write(target)
        key, value = "pdf:Title", "Original"
    elif kind == "docx":
        from docx import Document
        doc = Document()
        doc.core_properties.title = "Original"
        doc.save(target)
        key, value = "office:title", "Original"
    else:
        import mutagen.id3 as id3
        target.write_bytes(bytes([0xFF, 0xFB, 0x90, 0]) + bytes(413))
        with target.open("ab") as stream:
            stream.write((bytes([0xFF, 0xFB, 0x90, 0]) + bytes(413)) * 3)
        tags = id3.ID3()
        tags.add(id3.COMM(encoding=3, lang="eng", desc="custom", text=["Original"]))
        tags.save(target)
        key, value = "tag:COMM:custom:eng", "Original"
    response = await client.post("/delete", json={"path": str(target), "keys": [key]})
    assert response.status_code == 200, response.text
    response = await client.post("/write", json={"path": str(target), "fields": [
        {"key": key, "value": value, "editable": False, "source": "untrusted"},
    ]})
    assert response.status_code == 200, response.text
    response = await client.get("/read", params={"path": str(target)})
    restored = {field["key"]: field for field in response.json()["fields"]}
    assert restored[key]["value"] == value
    assert restored[key]["source"] != "untrusted"


@pytest.mark.parametrize("operation", ["write", "delete"])
def test_encrypted_pdf_cannot_lose_encryption(tmp_path, operation):
    from pypdf import PdfReader, PdfWriter
    from core.handlers.pdf_handler import PdfHandler
    target = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_metadata({"/Title": "Original"})
    writer.encrypt("")
    writer.write(target)
    before = target.read_bytes()
    handler = PdfHandler()
    assert not handler.read(target).supports_write
    with pytest.raises(ValueError, match="encrypted"):
        if operation == "write":
            handler.write(target, [MetadataField("pdf:Title", "Title", "changed")])
        else:
            handler.delete(target, ["pdf:Title"])
    assert target.read_bytes() == before
    assert PdfReader(target).is_encrypted


@pytest.mark.skipif(sys.platform != "win32", reason="NTFS named streams are Windows-specific")
def test_windows_atomic_write_preserves_named_streams(tmp_path):
    target = tmp_path / "file.txt"
    target.write_text("original")
    stream = Path(str(target) + ":private")
    stream.write_text("preserve")
    security.secure_atomic_write(target, lambda p: p.write_text("changed"))
    assert target.read_text() == "changed"
    assert stream.read_text() == "preserve"


@pytest.mark.skipif(sys.platform != "win32", reason="ReplaceFileW recovery is Windows-specific")
def test_windows_replace_recovers_original_after_partial_failure(tmp_path, monkeypatch):
    import os
    target = tmp_path / "file.txt"
    target.write_text("original")
    def partial_failure(original, replacement, backup):
        os.replace(original, backup)
        raise OSError("simulated ReplaceFileW error 1177")
    monkeypatch.setattr(security, "_windows_replace", partial_failure)
    with pytest.raises(OSError, match="1177"):
        security.secure_atomic_write(target, lambda p: p.write_text("changed"))
    assert target.read_text() == "original"
    assert list(tmp_path.iterdir()) == [target]


@pytest.mark.skipif(sys.platform != "win32", reason="DACL preservation is Windows-specific")
def test_windows_atomic_write_preserves_protected_dacl(tmp_path):
    import ntsecuritycon
    import win32security
    target = tmp_path / "file.txt"
    target.write_text("original")
    owner = win32security.GetFileSecurity(str(target), win32security.OWNER_SECURITY_INFORMATION).GetSecurityDescriptorOwner()
    acl = win32security.ACL()
    acl.AddAccessAllowedAce(win32security.ACL_REVISION, ntsecuritycon.FILE_ALL_ACCESS, owner)
    flags = win32security.DACL_SECURITY_INFORMATION | win32security.PROTECTED_DACL_SECURITY_INFORMATION
    win32security.SetNamedSecurityInfo(str(target), win32security.SE_FILE_OBJECT, flags, None, None, acl, None)
    def descriptor():
        security = win32security.GetFileSecurity(str(target), win32security.DACL_SECURITY_INFORMATION)
        return win32security.ConvertSecurityDescriptorToStringSecurityDescriptor(
            security, win32security.SDDL_REVISION_1, win32security.DACL_SECURITY_INFORMATION,
        )
    before = descriptor()
    security.secure_atomic_write(target, lambda p: p.write_text("changed"))
    assert descriptor() == before


def test_posix_owner_preservation_calls_chown_and_restores_mode(tmp_path, monkeypatch):
    from types import SimpleNamespace
    target = tmp_path / "file.txt"
    target.write_text("original")
    calls = []
    monkeypatch.setattr(security.os, "chown", lambda *args: calls.append(("chown", args)), raising=False)
    monkeypatch.setattr(security.os, "chmod", lambda *args, **kwargs: calls.append(("chmod", args)))
    original = SimpleNamespace(st_uid=12345, st_gid=23456, st_mode=0o100640)
    security._preserve_posix_owner(target, original)
    assert calls == [("chown", (target, 12345, 23456)), ("chmod", (target, 0o640))]


def test_posix_owner_preservation_refuses_permission_failure(tmp_path, monkeypatch):
    from types import SimpleNamespace
    target = tmp_path / "file.txt"
    target.write_text("original")
    def denied(*args):
        raise PermissionError("Cannot preserve ownership")
    monkeypatch.setattr(security.os, "chown", denied, raising=False)
    with pytest.raises(PermissionError, match="ownership"):
        security._preserve_posix_owner(target, SimpleNamespace(st_uid=12345, st_gid=23456, st_mode=0o100640))
    assert target.read_text() == "original"


@pytest.mark.skipif(sys.platform == "win32", reason="Real POSIX UID/GID preservation requires Unix")
def test_posix_atomic_write_preserves_real_owner(tmp_path):
    target = tmp_path / "file.txt"
    target.write_text("original")
    before = target.stat()
    security.secure_atomic_write(target, lambda p: p.write_text("changed"))
    after = target.stat()
    assert (after.st_uid, after.st_gid) == (before.st_uid, before.st_gid)
