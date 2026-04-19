from __future__ import annotations
from pathlib import Path
from typing import Any

from core.base_handler import BaseMetadataHandler
from core.models import MetadataField, MetadataRecord

_CORE_PROPS = [
    ("title", "Title"),
    ("author", "Author"),
    ("subject", "Subject"),
    ("keywords", "Keywords"),
    ("description", "Description"),
    ("last_modified_by", "Last Modified By"),
    ("revision", "Revision"),
    ("category", "Category"),
    ("content_status", "Content Status"),
    ("identifier", "Identifier"),
    ("language", "Language"),
    ("version", "Version"),
    ("created", "Created"),
    ("modified", "Modified"),
]


class OfficeHandler(BaseMetadataHandler):
    NAME = "office"
    EXTENSIONS = frozenset({".docx", ".xlsx", ".pptx", ".docm", ".xlsm", ".pptm"})
    PRIORITY = 15

    def read(self, path: Path) -> MetadataRecord:
        fields: list[MetadataField] = []
        errors: list[str] = []
        ext = path.suffix.lower()
        try:
            props = _load_core_props(ext, path)
            for attr, label in _CORE_PROPS:
                val = getattr(props, attr, None)
                if val is None:
                    continue
                if hasattr(val, "isoformat"):
                    val = val.isoformat()
                fields.append(MetadataField(
                    f"office:{attr}", label, str(val),
                    source="office"
                ))
        except Exception as e:
            errors.append(f"Office read error: {e}")

        record = self._make_record(path, read_errors=errors,
                                   supports_write=True, supports_delete=True)
        record.fields = fields
        return record

    def write(self, path: Path, fields: list[MetadataField]) -> None:
        def _do_write(tmp: Path) -> None:
            ext = tmp.suffix.lower()
            doc = _load_doc(ext, tmp)
            props = doc.core_properties
            for f in fields:
                if not f.key.startswith("office:"):
                    continue
                attr = f.key[len("office:"):]
                if hasattr(props, attr):
                    try:
                        setattr(props, attr, f.value)
                    except Exception:
                        pass
            doc.save(str(tmp))

        self._atomic_write(path, _do_write)

    def delete(self, path: Path, keys: list[str]) -> None:
        def _do_delete(tmp: Path) -> None:
            ext = tmp.suffix.lower()
            doc = _load_doc(ext, tmp)
            props = doc.core_properties
            for key in keys:
                attr = key[len("office:"):] if key.startswith("office:") else key
                if hasattr(props, attr):
                    try:
                        setattr(props, attr, None)
                    except Exception:
                        pass
            doc.save(str(tmp))

        self._atomic_write(path, _do_delete)


def _load_core_props(ext: str, path: Path):
    return _load_doc(ext, path).core_properties


def _load_doc(ext: str, path: Path):
    if ext in {".docx", ".docm"}:
        from docx import Document
        return Document(str(path))
    if ext in {".xlsx", ".xlsm"}:
        import openpyxl
        wb = openpyxl.load_workbook(str(path))
        wb.core_properties = wb.properties  # type: ignore
        return wb
    if ext in {".pptx", ".pptm"}:
        from pptx import Presentation
        return Presentation(str(path))
    raise ValueError(f"Unsupported extension: {ext}")
